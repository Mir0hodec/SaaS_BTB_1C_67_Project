"""MCP-сервер файлов одного вопроса.

- create_excel / create_word: модель передаёт содержимое, файл собирает этот
  сервер и кладёт в OUTBOX_DIR — мост отправит его сотруднику.
- list_attachments / view_attachment: вложения сотрудника из ATTACHMENTS_DIR.
  Картинки (скриншоты ошибок) отдаются модели как изображения, документы — текстом.
Доступа к остальной файловой системе нет.
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from common.win_utf8 import ensure_utf8_stdio  # noqa: E402

ensure_utf8_stdio()

from docx import Document  # noqa: E402
from mcp.server.mcpserver import Image, MCPServer  # noqa: E402
from mcp.server.mcpserver.exceptions import ToolError  # noqa: E402

from common.doc_extract import SUPPORTED_EXTENSIONS, extract_text  # noqa: E402
from common.xlsx import safe_file_name, write_workbook  # noqa: E402

OUTBOX = Path(os.environ.get("OUTBOX_DIR", ""))
ATTACHMENTS = Path(os.environ.get("ATTACHMENTS_DIR", ""))
IMAGE_EXT = {".png": "png", ".jpg": "jpeg", ".jpeg": "jpeg", ".gif": "gif", ".webp": "webp"}
MAX_TEXT_CHARS = 60_000

mcp = MCPServer(name="files", instructions="Создание Excel/Word для ответа и чтение вложений сотрудника.")


def _outbox() -> Path:
    if not str(OUTBOX):
        raise ToolError("Каталог для файлов не задан")
    OUTBOX.mkdir(parents=True, exist_ok=True)
    return OUTBOX


@mcp.tool()
def create_excel(file_name: str, sheets: list[dict]) -> str:
    """Создать xlsx. sheets: [{"name": "Лист", "columns": ["Колонка", ...],
    "rows": [[значение, ...], ...]}]. Даты — строкой ISO (ГГГГ-ММ-ДД), числа —
    числом: в файле будут настоящие даты и суммы, фильтры и закреплённая шапка."""
    name = safe_file_name(file_name, ".xlsx")
    rows = write_workbook(_outbox() / name, sheets)
    return f"Файл {name} готов ({rows} строк), будет отправлен сотруднику."


@mcp.tool()
def create_word(file_name: str, title: str, text: str) -> str:
    """Создать docx: заголовок + текст. Абзацы разделяй пустой строкой, строки
    с «- » станут маркированным списком. Подходит для ответа на письмо, справки."""
    name = safe_file_name(file_name, ".docx")
    doc = Document()
    if title:
        doc.add_heading(title, level=1)
    for block in text.replace("\r\n", "\n").split("\n\n"):
        lines = [ln for ln in block.split("\n") if ln.strip()]
        if lines and all(ln.lstrip().startswith(("- ", "• ")) for ln in lines):
            for ln in lines:
                doc.add_paragraph(ln.lstrip()[2:].strip(), style="List Bullet")
        elif lines:
            doc.add_paragraph("\n".join(lines))
    doc.save(_outbox() / name)
    return f"Файл {name} готов, будет отправлен сотруднику."


def _attachments() -> list[Path]:
    if not str(ATTACHMENTS) or not ATTACHMENTS.is_dir():
        return []
    return sorted(p for p in ATTACHMENTS.iterdir() if p.is_file())


@mcp.tool()
def list_attachments() -> list[dict]:
    """Вложения, которые сотрудник прислал к этому вопросу."""
    return [{"name": p.name, "size_kb": round(p.stat().st_size / 1024, 1),
             "kind": "image" if p.suffix.lower() in IMAGE_EXT else "document"} for p in _attachments()]


@mcp.tool()
def view_attachment(name: str):
    """Открыть вложение: картинку — как изображение, документ — текстом."""
    path = next((p for p in _attachments() if p.name == name), None)
    if path is None:
        raise ToolError(f"Вложения «{name}» нет — возьми имя из list_attachments")
    ext = path.suffix.lower()
    if ext in IMAGE_EXT:
        return Image(path=path, format=IMAGE_EXT[ext])
    if ext in SUPPORTED_EXTENSIONS:
        text = extract_text(path)
        return text[:MAX_TEXT_CHARS] + ("\n…(обрезано)" if len(text) > MAX_TEXT_CHARS else "")
    raise ToolError(f"Формат {ext} не читается: поддерживаются картинки и {', '.join(sorted(SUPPORTED_EXTENSIONS))}")


if __name__ == "__main__":
    mcp.run(transport="stdio")
