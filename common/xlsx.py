"""Запись таблиц в xlsx: фильтры, закреплённая шапка, настоящие даты и числа."""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

_ISO_DT_RE = re.compile(r"^\d{4}-\d{2}-\d{2}(T\d{2}:\d{2}(:\d{2})?)?")
_NUM_RE = re.compile(r"^-?\d+([.,]\d+)?$")
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _cell(value):
    if isinstance(value, str):
        v = value.strip()
        if _ISO_DT_RE.match(v):
            try:
                parsed = datetime.fromisoformat(v.replace("Z", "+00:00")).replace(tzinfo=None)
                return parsed.date() if len(v) == 10 else parsed
            except ValueError:
                return value
        if _NUM_RE.match(v) and not (len(v) > 1 and v.startswith("0") and v[1] not in ".,"):
            # «0012» — код/номер, а не число: ведущие нули не теряем.
            return float(v.replace(",", ".")) if any(c in v for c in ".,") else int(v)
    return value


def write_workbook(path: Path, sheets: list[dict]) -> int:
    """sheets: [{"name": str, "columns": [str], "rows": [[...]]}]. Возвращает число строк."""
    wb = Workbook()
    wb.remove(wb.active)
    total = 0
    for i, sheet in enumerate(sheets or [{"name": "Лист1", "columns": [], "rows": []}]):
        name = _BAD_SHEET_CHARS.sub(" ", str(sheet.get("name") or f"Лист{i + 1}"))[:31] or f"Лист{i + 1}"
        ws = wb.create_sheet(name)
        columns = [str(c) for c in sheet.get("columns") or []]
        rows = sheet.get("rows") or []
        if columns:
            ws.append(columns)
            for c in ws[1]:
                c.font = Font(bold=True)
        widths = [len(c) for c in columns]
        for row in rows:
            values = [_cell(v) for v in row]
            ws.append(values)
            for j, v in enumerate(values):
                text_len = len(v.strftime("%d.%m.%Y %H:%M")) if isinstance(v, datetime) else len(str(v))
                if j >= len(widths):
                    widths.append(0)
                widths[j] = max(widths[j], min(text_len, 60))
        for r in ws.iter_rows(min_row=2):
            for c in r:
                if isinstance(c.value, datetime):
                    c.number_format = "DD.MM.YYYY HH:MM" if (c.value.hour or c.value.minute) else "DD.MM.YYYY"
                elif isinstance(c.value, date):
                    c.number_format = "DD.MM.YYYY"
                elif isinstance(c.value, float):
                    c.number_format = "#,##0.00"
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = max(8, w + 2)
        if columns:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"
        total += len(rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return total


_SAFE_NAME_RE = re.compile(r"[^\w\s.,()№-]+", re.U)


def safe_file_name(name: str, ext: str) -> str:
    """Только имя файла, без путей и спецсимволов; расширение принудительно."""
    base = Path(str(name).replace("\\", "/")).name
    base = _SAFE_NAME_RE.sub("_", base).strip(" ._") or "файл"
    stem = base[: -len(ext)] if base.lower().endswith(ext) else Path(base).stem or base
    return f"{stem[:100]}{ext}"
