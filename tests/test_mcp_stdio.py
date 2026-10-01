"""MCP-серверы как отдельные процессы по stdio, по конфигу моста, из пустой папки."""
import asyncio
import base64
import json
import os
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

from mcp import ClientSession
from mcp.client.stdio import StdioServerParameters, stdio_client
from openpyxl import load_workbook

from bridge.config import Settings
from bridge.mcp_config_builder import build_mcp_config


async def _call(cfg: dict, cwd: Path, calls: list[tuple[str, dict]]):
    params = StdioServerParameters(command=cfg["command"], args=cfg["args"], env=cfg["env"], cwd=str(cwd))
    async with stdio_client(params) as (r, w), ClientSession(r, w) as s:
        await s.initialize()
        return [await s.call_tool(name, args) for name, args in calls]


def test_onec_and_files_servers(repo_root, tmp_path):
    settings = Settings.load(repo_root / "config" / "settings.example.yaml")
    settings.bases_path = repo_root / "config" / "bases.example.yaml"
    work, out, att = tmp_path / "work", tmp_path / "out", tmp_path / "att"
    work.mkdir()
    att.mkdir()
    (att / "письмо.txt").write_text("Прошу выслать акт сверки", encoding="utf-8")
    servers = json.loads(build_mcp_config(settings, base_id="demo_base", outbox=out, attachments=att,
                                          target=tmp_path / "mcp.json").read_text(encoding="utf-8"))["mcpServers"]

    meta, export, code, bad = asyncio.run(_call(servers["onec"], work, [
        ("get_metadata", {"object_name": "Документ.РасходнаяНакладная"}),
        ("export_query_to_excel", {
            "query_text": "ВЫБРАТЬ * ИЗ Документ.РасходнаяНакладная ГДЕ Дата МЕЖДУ &НачалоПериода И &КонецПериода",
            "date_from": "2026-01-15", "date_to": "2026-03-10", "file_name": "../../реестр"}),
        ("search_code", {"text": "ОстаткиТоваров"}),
        ("export_query_to_excel", {"query_text": "ВЫБРАТЬ 1", "date_from": "2026-01-01",
                                   "date_to": "2026-01-31", "file_name": "x"}),
    ]))
    assert "Контрагент" in meta.content[0].text
    info = json.loads(export.content[0].text)
    assert info["file"] == "реестр.xlsx" and info["months"] == 3 and info["rows"] == 90
    ws = load_workbook(out / "реестр.xlsx").active
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1")
    assert ws["A2"].is_date and isinstance(ws["D2"].value, float)
    assert json.loads(code.content[0].text)["total"] == 3
    assert bad.is_error and "НачалоПериода" in bad.content[0].text

    listing, text, word = asyncio.run(_call(servers["files"], work, [
        ("list_attachments", {}),
        ("view_attachment", {"name": "письмо.txt"}),
        ("create_word", {"file_name": "ответ", "title": "Ответ на письмо", "text": "Акт направим.\n\n- пункт 1"}),
    ]))
    assert "письмо.txt" in listing.content[0].text
    assert "акт сверки" in text.content[0].text
    assert not word.is_error and (out / "ответ.docx").exists()


def test_counterparty_server_over_stdio(repo_root, tmp_path, monkeypatch):
    seen = {}

    class FakeOneC(BaseHTTPRequestHandler):
        def do_POST(self):
            seen["auth"] = self.headers.get("Authorization")
            seen["body"] = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            data = json.dumps({"counterpartyName": "ТОО Альфа", "debtAmount": 1500,
                               "unsignedEavrCount": 1, "unsignedEavrAmount": 300}).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)

        def log_message(self, *args):
            pass

    server = HTTPServer(("127.0.0.1", 0), FakeOneC)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    monkeypatch.setenv("COUNTERPARTY_CHECK_URL", f"http://127.0.0.1:{server.server_port}/gke/hs/CounterpartyCheck/check")
    monkeypatch.setenv("COUNTERPARTY_CHECK_LOGIN", "svc")
    monkeypatch.setenv("COUNTERPARTY_CHECK_PASSWORD", "S3cr3t-P@ss")

    settings = Settings.load(repo_root / "config" / "settings.example.yaml")
    config_file = build_mcp_config(settings, base_id=None, outbox=tmp_path / "out", attachments=tmp_path / "att",
                                   target=tmp_path / "mcp.json")
    assert "S3cr3t-P@ss" not in config_file.read_text(encoding="utf-8")
    cfg = json.loads(config_file.read_text(encoding="utf-8"))["mcpServers"]["counterparty"]
    # Процесс claude передаёт серверу своё окружение — имитируем это наследование.
    cfg = {**cfg, "env": {**os.environ, **cfg["env"]}}
    work = tmp_path / "work"
    work.mkdir()
    try:
        (result,) = asyncio.run(_call(cfg, work, [("check_counterparty", {"bin": "750710450345"})]))
    finally:
        server.shutdown()

    data = json.loads(result.content[0].text)
    assert data["status"] == "needs_manager" and data["debt_amount"] == 1500.0
    assert "клиент-менеджеру" in data["decision"] and "raw_response" not in data
    assert seen["body"] == {"bin": "750710450345"}
    assert seen["auth"] == "Basic " + base64.b64encode(b"svc:S3cr3t-P@ss").decode()


def test_counterparty_tool_only_when_configured(repo_root, tmp_path, monkeypatch):
    from bridge.mcp_config_builder import allowed_tools
    monkeypatch.delenv("COUNTERPARTY_CHECK_URL", raising=False)
    settings = Settings.load(repo_root / "config" / "settings.example.yaml")
    cfg = json.loads(build_mcp_config(settings, base_id=None, outbox=tmp_path, attachments=tmp_path,
                                      target=tmp_path / "m.json").read_text(encoding="utf-8"))
    assert "counterparty" not in cfg["mcpServers"]
    assert "mcp__counterparty__check_counterparty" not in allowed_tools(None, settings.counterparty_check_enabled)
    assert "mcp__counterparty__check_counterparty" in allowed_tools(None, True)
